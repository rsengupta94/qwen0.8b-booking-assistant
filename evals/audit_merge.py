"""Score the simulator-fidelity audit: the human's blind category picks against the cards' intended category.

Per row: faithful (first pick matches the card), faithful_second (first pick matches, but the human also named
a second choice), ambiguous (only the second choice matches, or "can't tell"), unfaithful (a different category),
unlabelled (no pick). faithful_rate counts both faithful kinds. Card-verbatim rows are counted separately because
their wording was written into the card, not by the simulator.
Usage: uv run python -m evals.audit_merge
Prints counts, then every row that is not plainly faithful once every row is labelled. Writes nothing.
"""
from __future__ import annotations

import json
from collections import Counter

from evals.audit_sample import CANT_TELL, KEY, SHEET


def read_labels() -> dict[int, dict]:
    from openpyxl import load_workbook
    ws = load_workbook(SHEET, read_only=True)["labels"]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(it)]
    out = {}
    for r in it:
        if not r or r[0] is None:
            continue
        d = {k: ("" if v is None else str(v).strip()) for k, v in zip(hdr, r)}
        out[int(float(d["row"]))] = d
    return out


def bucket(card: str, first: str, second: str) -> str:
    if not first:
        return "unlabelled"
    if first == card:
        return "faithful_second" if second else "faithful"
    if first == CANT_TELL or second == card:
        return "ambiguous"
    return "unfaithful"


def main() -> None:
    labels = read_labels()
    key = json.loads(KEY.read_text())
    counts = {"simulator": Counter(), "card_verbatim": Counter()}
    flagged = []
    for k in key:
        lab = labels.get(k["row"], {})
        b = bucket(k["card_category"], lab.get("your_category", ""), lab.get("second_choice", ""))
        counts[k["source"]][b] += 1
        if b != "faithful":
            flagged.append((k, lab, b))
    for src, c in counts.items():
        match = c["faithful"] + c["faithful_second"]
        scored = match + c["ambiguous"] + c["unfaithful"]
        rate = round(match / scored, 3) if scored else None
        print(f"{src}: rows={sum(c.values())} faithful={c['faithful']} faithful_second={c['faithful_second']} "
              f"ambiguous={c['ambiguous']} unfaithful={c['unfaithful']} unlabelled={c['unlabelled']} faithful_rate={rate}")
    # Listing rows reveals the key, so it waits until every row has a pick.
    unlabelled = sum(c["unlabelled"] for c in counts.values())
    if unlabelled:
        print(f"{unlabelled} rows unlabelled; the row list prints once every row has a pick")
        return
    for k, lab, b in flagged:
        print(f"  row {k['row']} {b}: card={k['card_category']} you={lab.get('your_category') or '-'} "
              f"second={lab.get('second_choice') or '-'} [{k['card_id']} turn {k['turn']}, {k['source']}] "
              f"note: {lab.get('note') or '-'}")


if __name__ == "__main__":
    main()
