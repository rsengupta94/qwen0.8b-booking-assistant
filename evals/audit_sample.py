"""Build the blind simulator-fidelity audit sheet from dev transcripts (eval design section 7).

One row per patient message that describes the problem: the first answer at ASK_PROBLEM and any
later correction of the problem field. Rows are shuffled and carry no card id or category; the
card's intended category goes to a separate key file the labeller does not open.
Writes evals/calibration/audit.xlsx (sheets 'instructions', 'labels') and evals/calibration/audit_key.json.
Refuses to overwrite an existing sheet, so labels are never wiped. Dev pool only.
Usage: uv run python -m evals.audit_sample [--force]
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
SEED = 20261004  # changed from 20261003 after a dry run printed that order's key
CANT_TELL = "can't tell"
SHEET = ROOT / "calibration" / "audit.xlsx"
KEY = ROOT / "calibration" / "audit_key.json"

# Plain-English reading of each category name, for the labeller only. The product sees just the name.
GLOSS = {
    "addiction": "alcohol, drugs, smoking, gaming or phone overuse",
    "anxiety": "worry, panic, fear, nervousness",
    "child_adolescent": "the patient is a child or teenager",
    "depression": "low mood, loss of interest, hopelessness",
    "geriatric": "the patient is an older adult, problems tied to ageing",
    "grief": "loss of someone, bereavement",
    "ocd": "intrusive thoughts, compulsions, checking or cleaning rituals",
    "perinatal": "pregnancy, after childbirth, fertility",
    "relationships": "partner, marriage, family conflict",
    "sleep": "trouble falling or staying asleep, irregular sleep",
    "stress": "pressure from work, exams, caregiving; burnout",
    "trauma": "a frightening or violent event, flashbacks",
}

INSTRUCTIONS = [
    ("What this is", ""),
    ("", "A check on the simulated patient, not on the booking assistant. Each persona card fixed a problem "
         "category (for example grief), and the simulator had to describe it in everyday words without naming it. "
         "The eval then marks the assistant wrong whenever it picks a different category. If the simulator's "
         "description actually points somewhere else, that 'wrong' is the simulator's fault, and the eval is "
         "blaming the product for it. Your labels measure how often that happens."),
    ("", ""),
    ("What you see", ""),
    ("", "Sheet 'labels': one row per patient message that describes their problem, from the 20 dev conversations "
         "only. Rows are shuffled. The card and its intended category are hidden on purpose: knowing the answer "
         "first pulls your pick towards it."),
    ("", "Do not open evals/calibration/audit_key.json until you have finished labelling."),
    ("", ""),
    ("Columns", ""),
    ("row", "Row number. Do not edit."),
    ("message_type", "'first description' is the patient's answer to 'what would you like help with'. "
                     "'changed description' is the patient later correcting what they said earlier; judge it on its own."),
    ("patient_message", "Exactly what the simulated patient typed. This is what you judge."),
    ("your_category", "Required. Pick from the dropdown the one category you would file this message under, "
                      "or \"can't tell\" if the message gives you no basis to choose."),
    ("second_choice", "Optional. Fill only if a second category is genuinely as plausible as your first. "
                      "Leave blank if your first pick is clearly better."),
    ("note", "Optional. One line on why, most useful for a second choice or \"can't tell\"."),
    ("", ""),
    ("How to judge", ""),
    ("1", "Read the message as a clinic receptionist who knows only the 12 categories below and has to send the "
          "patient to the right kind of doctor. Use only what the message says."),
    ("2", "Ask: what is the main thing this person wants help with? Symptoms often overlap (poor sleep after a "
          "death, stress that shows up as panic). Pick the cause or focus the patient puts first, not every "
          "symptom they mention."),
    ("3", "Ignore anything in the message that is not about the problem: questions about fees, 'is this a bot', "
          "days or doctor names."),
    ("4", "Do not guess what the scenario was meant to test, and do not look at transcripts, cards or the product."),
    ("5", "Use second_choice when you are truly torn between two. Use \"can't tell\" when the message is too vague "
          "for any category. Both are useful answers, not failures."),
    ("6", "About a minute per row. Save the file in place, same name, without renaming the sheets."),
    ("", ""),
    ("What happens next", ""),
    ("", "Run: uv run python -m evals.audit_merge. It compares your picks with the hidden key and reports three "
         "counts: faithful (your first pick matches the card), ambiguous (only your second choice matches, or you "
         "said can't tell), and unfaithful (you picked a different category). Rows whose wording comes from the "
         "card itself (a fixed sentence written by you, not the simulator) are reported separately."),
    ("", ""),
    ("Categories", ""),
] + [(c, g) for c, g in GLOSS.items()]


def categories() -> list[str]:
    """Same derivation as the product, read from the fixture file; eval code never imports app/."""
    doctors = json.loads((REPO / "fixtures" / "doctors.json").read_text())
    return sorted({c for d in doctors for c in d["categories"]})


def rows_from(t: dict, card: dict) -> list[dict]:
    verbatim = (card.get("verbatim") or {}).get("ASK_PROBLEM")
    out = []
    for tr in t["turns"]:
        s = tr["sidecar"]
        if s.get("field") != "problem" or s["act"] not in ("answer", "correction"):
            continue
        first = s["act"] == "answer"
        out.append({"card_id": t["card_id"], "turn": tr["turn"], "card_category": s["value"],
                    "message_type": "first description" if first else "changed description",
                    "source": "card_verbatim" if (first and verbatim and verbatim in tr["user"]) else "simulator",
                    "patient_message": tr["user"]})
    return out


def write_sheet(rows: list[dict], cats: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ins = wb.active
    ins.title = "instructions"
    ins.column_dimensions["A"].width = 20
    ins.column_dimensions["B"].width = 110
    for i, (a, b) in enumerate(INSTRUCTIONS, start=1):
        ins.cell(i, 1, a).font = Font(bold=bool(a) and not b)
        ins.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
        ins.cell(i, 1).alignment = Alignment(vertical="top")

    ws = wb.create_sheet("labels")
    hdr = ["row", "message_type", "patient_message", "your_category", "second_choice", "note"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["row"], r["message_type"], r["patient_message"], None, None, None])
    for col, w in zip("ABCDEF", (6, 20, 80, 20, 20, 40)):
        ws.column_dimensions[col].width = w
    for line in ws.iter_rows(min_row=2):
        for c in line:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    pick = DataValidation(type="list", formula1='"' + ",".join(cats + [CANT_TELL]) + '"', allow_blank=True)
    second = DataValidation(type="list", formula1='"' + ",".join(cats) + '"', allow_blank=True)
    ws.add_data_validation(pick)
    ws.add_data_validation(second)
    pick.add(f"D2:D{len(rows) + 1}")
    second.add(f"E2:E{len(rows) + 1}")
    wb.active = 0
    wb.save(SHEET)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="overwrite an existing sheet (wipes its labels)")
    a = ap.parse_args()
    if SHEET.exists() and not a.force:
        sys.exit(f"refusing: {SHEET.relative_to(REPO)} exists and may hold labels; pass --force to overwrite")
    rows = []
    for p in sorted((ROOT / "transcripts" / "dev").glob("*.json")):
        t = json.loads(p.read_text())
        card = json.loads((ROOT / "cards" / "dev" / p.name).read_text())
        rows.extend(rows_from(t, card))
    random.Random(SEED).shuffle(rows)
    for i, r in enumerate(rows, start=1):
        r["row"] = i
    cats = categories()
    missing = set(GLOSS) ^ set(cats)
    if missing:
        sys.exit(f"category list changed, update GLOSS: {sorted(missing)}")
    write_sheet(rows, cats)
    KEY.write_text(json.dumps([{k: r[k] for k in ("row", "card_id", "turn", "card_category", "message_type", "source")}
                               for r in rows], indent=2) + "\n")
    n_verb = sum(r["source"] == "card_verbatim" for r in rows)
    print(f"wrote {len(rows)} rows ({n_verb} card-verbatim) to {SHEET.relative_to(REPO)}; key at {KEY.relative_to(REPO)}")


if __name__ == "__main__":
    main()
